"""Quantitative data: Databook draft parsing, coverage against Sustainability 2030 indicators,
the master's Quant tabs and each central unit's '4. Quantitative data' request tab (SR25 'CFOG - Quant.' style)."""
import difflib
import re
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")
SRC = Path(__file__).parent / "source"
YEARS = ["Baseline", "2023", "2024", "2025", "2026"]
SKIP_SHEETS = {"Cover", "Overview and Contents", "References", "Estate & Infrastructure"}  # E&I tab is hidden, all 'TBC'

# Suggested data owner by SR26 ref prefix (longest match wins). None = SST compiles, no request.
OWNER = {"EE": None, "EE1(c)": "SaSS", "EE1(d)": "SaSS", "TR": "MRE", "CL1": "ESG", "CL2": None, "CL3": "CI&S", "NB1": "ESG", "NB2": "MRE",
         "AI": "CIOG", "EI": "CI&S", "CE": "ESG", "RP": "CFOG", "RI": "CFOG", "EN": "CI&S"}
OWNER_NOTE = {"EE": "SST",
              "CL2": "SST – compiled from faculty and CI&S responses"}


def norm(t):
    return re.sub(r"[^a-z0-9 ]", " ", str(t).lower())


def owner_for(ref):
    best = max((p for p in OWNER if ref.startswith(p)), key=len, default=None)
    return (OWNER[best], OWNER_NOTE.get(best, "")) if best else (None, "")


def match_ref(text, qual, pa_rows):
    """Best S2030 Qual indicator for a Databook indicator text (restricted to the sheet's priority area)."""
    pool = pa_rows or qual
    scored = [(difflib.SequenceMatcher(None, norm(text)[:160], norm(q["indicator"])[:160]).ratio(), q) for q in pool]
    s, q = max(scored, key=lambda x: x[0])
    return q if s >= 0.45 else None


def read_databook(qual):
    """[{sheet, pa, ref, indicator, metric, level, unit, values{year}, has_data, group}] – one per Databook row."""
    wb = load_workbook(SRC / "Databook_draft.xlsx", data_only=True)
    out = []
    for ws in wb:
        if ws.title in SKIP_SHEETS:
            continue
        pa = ws.title.replace("&", "and")
        pa_rows = [q for q in qual if norm(q["pa"]).split()[:2] == norm(pa).split()[:2]]
        hdr = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 3).value or "").strip() == "Metrics"]
        indicator = ""
        for i, h in enumerate(hdr):
            above = str(ws.cell(h - 1, 3).value or "").strip()
            if above:
                indicator = above
            cols = {}
            for c in range(3, ws.max_column + 1):
                v = str(ws.cell(h, c).value or "").strip()
                if v == "Unit":
                    cols["unit"] = c
                elif v in YEARS or v[:4] in YEARS:
                    cols[v[:4] if v[:4].isdigit() else v] = c
            if "unit" not in cols:  # Nature sheet: unit column one left of 'Baseline'
                cols["unit"] = cols.get("Baseline", 8) - 1
            end = hdr[i + 1] - 1 if i + 1 < len(hdr) else ws.max_row
            q = match_ref(indicator, qual, pa_rows)
            rows = []
            for r in range(h + 1, end + 1):
                lab, lvl = "", 0
                for k, c in enumerate((3, 4, 5)):
                    v = ws.cell(r, c).value
                    if v not in (None, "") and str(v).strip() not in ("`",):
                        lab, lvl = str(v).strip(), k
                        break
                if not lab:
                    continue
                unit = str(ws.cell(r, cols["unit"]).value or "").strip()
                vals = {y: ws.cell(r, cols[y]).value for y in YEARS if y in cols}
                nums = any(isinstance(v, (int, float)) for v in vals.values())
                rows.append(dict(r=r, lab=lab, lvl=lvl, unit=unit, vals=vals, nums=nums))
            # drop trailing text rows (next indicator's target/indicator text, footnotes)
            keep, glvl = [], None
            for j, x in enumerate(rows):
                nxt = rows[j + 1] if j + 1 < len(rows) else None
                is_group = not x["unit"] and not x["nums"] and nxt and nxt["lvl"] > x["lvl"]
                in_group = glvl is not None and x["lvl"] > glvl and len(x["lab"]) < 90
                if is_group:
                    glvl = x["lvl"]
                elif glvl is not None and x["lvl"] <= glvl:
                    glvl = None
                if x["unit"] or x["nums"] or is_group or in_group:
                    keep.append(dict(x, group=bool(is_group)))
            for x in keep:
                out.append(dict(sheet=ws.title, pa=pa, ref=q["ref"] if q else "", qrow=q["row"] if q else None,
                                indicator=indicator, metric=x["lab"], level=x["lvl"], unit=x["unit"], values=x["vals"],
                                has_data=x["nums"], group=x["group"]))
    return out


# ---------------------------------------------------------------- classification
def quant_only_rows(qual):
    """Qual-sheet rows that are purely quantitative: no stakeholder named (RP1(a) case studies excepted – routed to CFOG)."""
    return {q["row"] for q in qual if not q["stake"] and q["row"] != 46}


def coverage(qual, quan, db):
    """Indicator-level rows for '2. Quant coverage'."""
    purged = quant_only_rows(qual)
    by_row = {q["row"]: q for q in qual}
    qn = {x["row"]: x for x in quan}
    db_by_ref = {}
    for d in db:
        if d["ref"]:
            db_by_ref.setdefault(d["ref"], []).append(d)
    rows = sorted(set(qn) | purged | {d["qrow"] for d in db if d["qrow"]})
    out = []
    for r in rows:
        q = by_row.get(r)
        if not q:
            continue
        pts = [d for d in db_by_ref.get(q["ref"], []) if not d["group"]]
        with_data = [d for d in pts if d["has_data"]]
        yrs = sorted({y for d in with_data for y, v in d["values"].items() if isinstance(v, (int, float)) and y.isdigit()})
        if with_data:
            indb = f"Yes – data held ({yrs[0]}–{yrs[-1]})" if len(yrs) > 1 else f"Yes – data held ({yrs[0] if yrs else 'baseline'})"
        elif pts:
            indb = "Yes – metrics defined, no data yet"
        else:
            indb = "No"
        own, note = owner_for(q["ref"])
        if own is None:
            action = note or "SST compiles – no request"
        elif with_data:
            action = "Owner confirms existing figures and provides CY2026 (provisional Jan, final Feb) – Tab 4 of owner's request"
        elif pts:
            action = "Owner provides CY2026 (and earlier years if available) and confirms metric definition – Tab 4 of owner's request"
        else:
            action = "Not in Databook – decide: add metric to Databook, or report this indicator qualitatively only"
        kind = ("Quantitative only – removed from requirements matrix" if r in purged
                else "Mixed – narrative in requirements matrix; metric here")
        out.append(dict(row=r, ref=q["ref"], pa=q["pa"], indicator=q["indicator"], kind=kind,
                        flag=qn.get(r, {}).get("flag", ""), metric=qn.get(r, {}).get("metric", "") or q["request"],
                        indb=indb, npts=len(pts), owner=own or "SST", action=action,
                        tab4=f"{own} request – Tab 4" if own and pts else ""))
    return out


def owned_points(db, code):
    return [d for d in db if d["ref"] and owner_for(d["ref"])[0] == code]


# ---------------------------------------------------------------- sheets
def coverage_tab(wb, cov):
    import build as B
    from openpyxl.styles import Font, PatternFill
    ws = wb.create_sheet("2. Quant coverage")
    B.title(ws, "End-2026 target status assessment and reporting", "2. Quantitative indicators – Databook coverage")
    ws["B3"] = ("The Sustainability Databook is the authoritative quantitative source (SR26 approach). Each quantitative indicator is checked "
                "against the Databook draft: figures held are confirmed by the owner; gaps are requested in Tab 4 of that owner's request. "
                "Faculties are not asked for quantitative data. Data points are listed in '2b. Databook register'.")
    ws["B3"].alignment = B.WRAP; ws.merge_cells("B3:M3"); ws.row_dimensions[3].height = 45
    cols = ["Ref", "Priority area", "Indicator", "Type", "Quant? (Y/M)", "Suggested metric (Quan sheet)", "In Databook draft?",
            "No. of data points", "Suggested data owner", "Action", "Requested in", "Status", "Notes"]
    B.hdr_row(ws, 5, cols)
    fills = {"Yes – data": "E2EFDA", "Yes – metrics": "FFF2CC", "No": "F8CBAD"}
    for k, c in enumerate(cov):
        rr = 6 + k
        vals = [c["ref"], c["pa"], c["indicator"], c["kind"], c["flag"], c["metric"], c["indb"], c["npts"], c["owner"],
                c["action"], c["tab4"], "Not started", ""]
        for i, v in enumerate(vals):
            B.body(ws.cell(rr, 2 + i, v), B.REFF if i == 0 else None, bold=i == 0)
        f = next((v for kk, v in fills.items() if c["indb"].startswith(kk)), None)
        if f:
            ws.cell(rr, 8).fill = PatternFill("solid", fgColor=f)
        ws.row_dimensions[rr].height = 75
    last = 5 + len(cov)
    B.dv_list(ws, '"Not started,Requested,Received,Confirmed,Not required"', f"M6:M{last}")
    B.widths(ws, dict(zip("BCDEFGHIJKLMN", [9, 18, 42, 24, 9, 40, 20, 10, 14, 46, 18, 13, 30])))
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"B5:N{last}"


def register_tab(wb, db):
    import build as B
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet("2b. Databook register")
    B.title(ws, "End-2026 target status assessment and reporting", "2b. Databook register (data points)")
    ws["B3"] = ("Every data point in the Databook draft. Baseline–CY2025 are copied from the draft for context. Complete the yellow columns as "
                "owners confirm figures (SR26 approach: definition, boundary, period, source, owner, quality note and confirmation for each figure).")
    ws["B3"].alignment = B.WRAP; ws.merge_cells("B3:N3"); ws.row_dimensions[3].height = 32
    cols = ["Ref", "Priority area", "Indicator", "Metric", "Unit", "Baseline", "CY2023", "CY2024", "CY2025", "CY2026",
            "Data status", "Data owner", "Definition", "Reporting boundary", "Source system", "Quality note / limitations",
            "Owner confirmed?", "Confirmation date"]
    B.hdr_row(ws, 5, cols)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    r, prev_ind = 6, None
    for d in db:
        own, note = owner_for(d["ref"]) if d["ref"] else (None, "")
        new_ind = d["indicator"] != prev_ind
        prev_ind = d["indicator"]
        v = d["values"]
        status = "" if d["group"] else ("Data held" if d["has_data"] else "No data yet")
        vals = [d["ref"], d["pa"] if new_ind else "", d["indicator"] if new_ind else "", d["metric"], d["unit"],
                v.get("Baseline"), v.get("2023"), v.get("2024"), v.get("2025"), None, status,
                "" if d["group"] else (own or note.split(" –")[0] or "SST")] + [""] * 6
        for i, x in enumerate(vals):
            c = ws.cell(r, 2 + i, x)
            B.body(c, B.GREY if d["group"] else (yellow if i >= 9 and i != 10 and i != 11 else None), bold=d["group"] or i == 0)
        ws.cell(r, 5).alignment = Alignment(indent=d["level"] * 2, wrap_text=True, vertical="top")
        for i in range(5, 10):
            ws.cell(r, 2 + i).number_format = "#,##0.##"
        r += 1
    B.dv_list(ws, "=Lists!$E$2:$E$4", f"R6:R{r}")
    B.widths(ws, dict(zip("BCDEFGHIJKLMNOPQRS", [9, 16, 34, 40, 12, 11, 11, 11, 11, 11, 12, 10, 22, 18, 18, 24, 11, 13])))
    ws.freeze_panes = "F6"
    ws.auto_filter.ref = f"B5:S{r - 1}"


def unit_quant_tab(wb, code, name, pts, due):
    """SR25 'CFOG - Quant.' style request tab. Returns set of refs covered."""
    import build as B
    from openpyxl.styles import Alignment, Font, PatternFill
    if not pts:
        return set()
    ws = wb.create_sheet("4. Quantitative data")
    B.title(ws, "End-2026 target status assessment and reporting", f"Quantitative data – {name}")
    ws["B3"] = ("These figures feed the Sustainability Databook (the University's authoritative quantitative source). Previous years are shown "
                "for context – please do not alter. Please provide the CY2026 value (provisional figures are fine; we will confirm final "
                f"figures with you in February), the data source, and any change to definition or boundary. Due: {due}.")
    ws["B3"].alignment = B.WRAP; ws.merge_cells("B3:O3"); ws.row_dimensions[3].height = 48
    gold = PatternFill("solid", fgColor="FFC000")
    ws.merge_cells("B5:J5"); ws["B5"] = "Context only (please do not alter)"
    ws.merge_cells("K5:O5"); ws["K5"] = "Please fill in these cells"
    for c, f, fc in (("B5", B.HDR, "FFFFFF"), ("K5", gold, B.DARK)):
        ws[c].fill = f; ws[c].font = Font(name="Aptos Narrow", size=14, bold=True, color=fc); ws[c].alignment = B.CWRAP
    cols = ["Ref", "Indicator", "Metric", "Unit", "Baseline", "CY2023", "CY2024", "CY2025", "Data status",
            "CY2026 value", "Data source (file name, link or system)", "Change to definition or boundary?", "Notes / limitations",
            "Useful case studies"]
    B.hdr_row(ws, 6, cols)
    for i in range(9, len(cols)):
        ws.cell(6, 2 + i).fill = gold; ws.cell(6, 2 + i).font = Font(name="Aptos", size=12, bold=True, color=B.DARK)
    tips = ["No instructions – context only.", "", "", "", "", "", "", "", "",
            "Please enter numerical values", "Please provide a file name or link to the data source",
            "Describe any change since last year (or 'No change')", "Please add any additional notes here",
            "Please provide information on any useful case studies"]
    for i, t in enumerate(tips):
        B.body(ws.cell(7, 2 + i, t), B.GREY); ws.cell(7, 2 + i).font = Font(name="Aptos Narrow", size=9, italic=True)
    r, prev = 8, None
    for d in pts:
        v = d["values"]
        vals = [d["ref"], d["indicator"] if d["indicator"] != prev else "", d["metric"], d["unit"], v.get("Baseline"),
                v.get("2023"), v.get("2024"), v.get("2025"), "" if d["group"] else ("Data held" if d["has_data"] else "New – no data yet")]
        prev = d["indicator"]
        for i, x in enumerate(vals):
            B.body(ws.cell(r, 2 + i, x), B.GREY if d["group"] else B.REFF if i == 0 else None, bold=d["group"] or i == 0)
        for i in range(9, len(cols)):
            B.body(ws.cell(r, 2 + i), B.GREY2 if d["group"] else None)
        ws.cell(r, 4).alignment = Alignment(indent=d["level"] * 2, wrap_text=True, vertical="top")
        for i in range(4, 8):
            ws.cell(r, 2 + i).number_format = "#,##0.##"
        r += 1
    B.dv_list(ws, '"No change,Yes – see notes"', f"M8:M{r}")
    B.widths(ws, dict(zip("BCDEFGHIJKLMNO", [9, 34, 40, 12, 11, 11, 11, 11, 14, 14, 30, 18, 30, 26])))
    ws.freeze_panes = "E8"
    return {d["ref"] for d in pts}
