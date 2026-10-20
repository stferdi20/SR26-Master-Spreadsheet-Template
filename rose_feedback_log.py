"""Rose feedback log (rounds 1 and 2): feedback, response and where to check. Run after build.py."""
import build as B
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
V = B.VERSION
M = f"Master ({V}), "
ABP = "ABP request, tab 1, "
MRE = "MRE request, tab 1, "
rows = [
 # round, file, Rose's feedback, response, applies to, where to check
 ("1", "ABP request", "D8: update the deadline, as the Sustainability 2030 launch was pushed back – Christmas shutdown is the deadline for responses.",
  "Due date changed to Thursday 24 December 2026 (Christmas shutdown). First review, follow-up and consolidation moved to January 2027.",
  "All 20 request files, master, project management file",
  ABP + "D4 (instructions, 'Due:' line). Master: 0. Read me (Key dates) and 4. Request tracker (Due date). Project file: P2.2 Timeline"),
 ("1", "ABP request", "F16 and all request cells with the same sentence: delete 'Please build on this rather than repeating it…' and move the end-2025 information to the bottom of the cell.",
  "Sentence removed. The end-2025 reference now sits at the end of every request cell.",
  "All 20 request files and master unit tabs",
  ABP + "F16 to F23 (last paragraph of each cell)"),
 ("1", "ABP request", "F16: add the note about the joint Associate Deans Sustainability and Associate Deans Teaching and Learning session on 23 November 2026.",
  "Rose's sentence added word for word to EE1(a).",
  "All 9 faculty requests", ABP + "F16 (EE1(a))"),
 ("1", "ABP request", "F16: delete the sentence on Melbourne Plus and the Wattle Fellowship.",
  "Sentence removed.", "All 9 faculty requests", ABP + "F16 (EE1(a))"),
 ("1", "ABP request", "F17: delete row (EE1(b) student sustainability self-efficacy).",
  "EE1(b) removed for ABP, then for all faculties (team decision). Faculties now have 8 questions. EE1(b) is not asked of faculties; logged in the Issues log.",
  "All 9 faculty requests, master Faculties tab, RASCI",
  ABP + "rows 16 to 23 (no EE1(b) row). " + M + "Issues log"),
 ("1", "ABP request", "F18 and F23: update CLLAP wording to 'Please note, we will collect case studies on Campus Living Lab Accelerator Program projects directly from project owners - no need to include these.'",
  "Rose's sentence used word for word in TR1(a) and EN1(b).",
  "All 9 faculty requests", ABP + "F17 (TR1(a)) and F22 (EN1(b))"),
 ("1", "ABP request", "F19: spell out MRE (Melbourne Research and Enterprise) as it is a new name.",
  "Written as 'Melbourne Research and Enterprise (MRE)' at first mention in request text, MRE file title and section names.",
  "All files", ABP + "F18 (TR2(a)). MRE request, tab 1, row 11"),
 ("1", "Project management", "P2 Gantt chart is not clear – format so blocks of time and dates are clear at a glance (refer to previous project plan).",
  "P2 split into three linked sheets. P2.2 Timeline: weekly grid with month headers, bars fill automatically from Start and End, dates written in the 'When' column, milestones shown as a diamond, Christmas shutdown greyed, committee labels above the weeks (SR25 style), activity columns frozen.",
  "Project management file", "Project file: P2.2 Timeline"),
 ("1", "Project management", "Note in the timeline when papers are due for governance committees (last year's dates as placeholders) and work backwards (+3 days leader review) to set internal deadlines.",
  "P2.1 Governance deadlines table: meeting, papers due, 'To leaders for review' (papers due minus 3 working days) and 'Begin drafting (proposed)' (minus 10 working days). Source column shows the SR26 approach date and last year's actual date. Changing a date updates P2.2.",
  "Project management file", "Project file: P2.1 Governance deadlines, rows 6 to 11"),
 ("1", "All spreadsheets", "Retain the formatting and look of the previous reference spreadsheets (Melbourne Uni blue band and logo at the top).",
  "Navy University of Melbourne banner with the logo (taken from the SR25 master) added to the top of every visible tab.",
  "All 22 files", "Rows 1 and 2 of any tab"),
 ("2", "MRE request", "MRE prepared all the indicators, so ask about every indicator, listed a to d under each target.",
  "MRE now has 12 rows, one per research indicator, in order TR1(a) to TR3(d).",
  "MRE request, master MRE tab", MRE + "rows 16 to 27. " + M + "MRE tab"),
 ("2", "MRE request", "F19: ask MRE to prepare the output from SciVal and Altmetric (Shane Wood has done this previously).",
  "TR2(c) request now reads: 'Please prepare the SciVal and Altmetric output on sustainability-related research…, as previously prepared by Shane Wood.'",
  "MRE request", MRE + "F22 (TR2(c))"),
 ("2", "MRE request", "F19: the cell refers to Tab 4 Quantitative data but the tab is missing – is it planned?",
  "Tab 4 Quantitative data is in the MRE file (Rose was likely viewing an older copy). It holds the research metrics MRE owns (TR2(b), TR2(c), TR3(b)).",
  "MRE request", "MRE request, tab '4. Quantitative data'"),
 ("2", "MRE request", "Rows 20 to 22: research target 3 is MRE's responsibility University-wide – ask about progress within MRE and University-wide.",
  "Each TR3 row opens with 'MRE is responsible for this target at a University-wide level – please describe progress within MRE and progress across the University.'",
  "MRE request", MRE + "F24 to F27 (TR3(a) to (d))"),
 ("2", "Master", "4. Request tracker: sheets are protected so column H cannot be expanded. Where is the source text, so changes made in the master flow to the data owner spreadsheets?",
  "The current master is not protected. Wording changes are made in the request files themselves (Review, Unprotect Sheet, no password). The master's Column F now pulls the wording from each request file, so edits show in the master after the links refresh.",
  "Master and all request files", M + "any unit tab, Column F. " + M + "Issues log (handover note)"),
 ("2", "Master – CI&S", "F10: we did receive an end-2025 response via another process – amend to '2025 response provided separately by Manager, Sustainability Strategy' and move to the end of the cell.",
  "Sentence replaced and placed at the end of the cell.", "CI&S request and master", M + "CI&S tab F10 (CL3(a)). CI&S request, tab 1, F16"),
 ("2", "Master – CI&S", "F11: indicator not similar enough to the one referenced, and no response last year – do not mention it.",
  "End-2025 reference removed from EI1(a).", "CI&S request and master", M + "CI&S tab F11 (EI1(a)). CI&S request, tab 1, F17"),
 ("2", "Master – CI&S", "G25: spell out Campus Living Lab Accelerator Program (CLLAP); don't say we will get the data from CI&S.",
  "Spelled out everywhere in request and 'How SST will use this data' text; now says case studies are collected directly from project owners.",
  "All files", M + "CI&S tab G25 (TR1(a)). CI&S request, tab 1, G31"),
 ("2", "Master – CI&S", "F28: the reporting request is the 2026 Greenhouse Gas Inventory.",
  "CL1(c) request now reads 'Please provide the 2026 Greenhouse Gas Inventory.'", "CI&S request and master",
  M + "CI&S tab F28 (CL1(c)). CI&S request, tab 1, F34"),
 ("2", "Master – CI&S", "F31: we don't report on procurement or modern slavery – no need to mention not receiving reporting last year.",
  "'We did not receive…' sentence removed from MS1(a).", "CI&S request and master", M + "CI&S tab F31 (MS1(a)). CI&S request, tab 1, F37"),
 ("2", "Master – CI&S", "F32: we don't report on investments – no need for the 'we didn't receive' message.",
  "'We did not receive…' sentence removed from RI1(b).", "CI&S request and master", M + "CI&S tab F32 (RI1(b)). CI&S request, tab 1, F38"),
 ("2", "Master – CFOG", "Delete TR3(b).", "TR3(b) removed from CFOG (MRE covers TR3 University-wide).", "CFOG request and master",
  M + "CFOG tab rows 10 to 14 (no TR3(b)). " + M + "Issues log"),
 ("2", "Master – CFOG", "F14 and all EN1(c) references: use 'sustainability outcomes' instead of 'Sustainability 2030 outcomes', as the strategy has only just been released.",
  "Changed in every EN1(c) request and in the 'How SST will use this data' text.", "All units asked EN1(c), incl. all faculties",
  M + "CFOG tab F13. ABP request, tab 1, F23"),
 ("2", "Master – AC&M", "AC&M's remit is University-wide – change 'your faculty/portfolio' to 'the University'.",
  "AC&M (and MRE) wording now refers to 'the University' / 'across the University'.", "AC&M and MRE requests",
  M + "AC&M tab F10 to F13. AC&M request (ACM file), tab 1, F16 to F19"),
 ("2", "Master – AC&M", "AC&M should also include the alumni profiles request (last indicator under Exceptional Education).",
  "EE1(d) alumni request added to AC&M; SaSS is still asked as well.", "AC&M request and master",
  M + "AC&M tab F10 (EE1(d)). ACM request, tab 1, F16"),
 ("2", "Master – AC&M", "Delete row 11 (CL2(a)).", "CL2(a) removed. No unit is now asked CL2(a); logged in the Issues log to decide later.",
  "AC&M request and master", M + "AC&M tab (no CL2(a) row). " + M + "Issues log"),
 ("2", "Master – other tabs", "Discuss how and where to update cells that say 'Reporting request TBC' during handover.",
  "Fill them in the request files themselves (Review, Unprotect Sheet); the master picks the wording up through Column F. Logged in the Issues log.",
  "Request files with TBC requests", M + "Issues log (handover note)"),
 ("2", "All data spreadsheets", "Columns P and Q: no need for a Yes/No 'Data owner confirms' column – the data owner is the most senior person, so comms will ask for reviewed and confirmed responses. Column O is enough to note people closer to the data.",
  "'Data owner confirms' and 'Confirmed by' removed. Column O renamed 'Data source, limitations and contact person'. Instructions ask for the file to be returned once the senior lead has reviewed and confirmed it. Tracker and response summary no longer count confirmations.",
  "All 20 request files and master", "ABP request, tab 1, row 14 (headers end at O) and D4. " + M + "4. Request tracker, A1. Response summary"),
 ("2", "All data spreadsheets", "Use 'faculty' for faculties and 'portfolio' for anything that isn't a faculty in Column F (only if possible).",
  "Done – wording is set per file: 'your faculty' for faculties, 'your portfolio' for central units, 'the University' for AC&M and MRE.",
  "All 20 request files", ABP + "F16 to F23. CFOG request, tab 1, F16 to F20"),
 ("Team", "MRE request", "(Team follow-up) NB2(a) nature footprint was included in MRE for the Melbourne Biodiversity Institute (MBI).",
  "MBI skipped for now – NB2(a) removed from the MRE request and logged in the Issues log.", "MRE request", MRE + "rows 16 to 27 (no NB2(a)). " + M + "Issues log"),
]
wb = Workbook(); ws = wb.active; ws.title = "Rose feedback log"
B.title(ws, "", "SR26 – Rose's feedback and responses")
ws["B3"] = f"Rounds 1 and 2 of Rose's review, what was changed, and where to check it. Files: {V} master, request files (no version in name), SR26 project management {V}."
ws["B3"].font = Font(name="Aptos", italic=True)
cols = ["#", "Round", "File / area", "Rose's feedback", "Response – what we did", "Applies to", "Where to check", "Status"]
B.hdr_row(ws, 5, cols)
for k, r in enumerate(rows):
    rr = 6 + k
    for i, v in enumerate([k + 1] + list(r[:4]) + [r[5], r[4]][::-1][::-1] + ["Done" if r[0] != "Team" or True else ""]):
        pass
    vals = [k + 1, r[0], r[1], r[2], r[3], r[4], r[5], "Done"]
    for i, v in enumerate(vals):
        c = ws.cell(rr, 2 + i, v); B.body(c, B.REFF if i in (0, 1, 2) else None)
B.dv_list(ws, '"Done,In progress,Open,Not actioned"', f"I6:I{5 + len(rows)}")
ws.conditional_formatting.add(f"I6:I{5 + len(rows)}", __import__("openpyxl").formatting.rule.CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="E2EFDA")))
B.widths(ws, {"A": 3, "B": 5, "C": 8, "D": 18, "E": 55, "F": 55, "G": 26, "H": 45, "I": 11})
ws.freeze_panes = "A6"
ws.auto_filter.ref = f"B5:I{5 + len(rows)}"
B.finalize(wb)
out = str(B.OUT / f"SR26 Rose feedback log - {V}.xlsx")
wb.save(out); print(out, len(rows))
