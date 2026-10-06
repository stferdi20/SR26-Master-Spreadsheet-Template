"""SR26 RASCI matrix, rebuilt to the SR25 'RASCI' tab standard (Sustainability Plan 2030 Implementation Roles):
Domain > Priority area > Target > Indicator hierarchy, consolidated portfolio columns with grouped sub-units,
SST contact per area, A / R / S / C / I codes, colour coding and COUNTIF summary rows."""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# SST contact per area, carried from the SR25 RASCI (row 3) – to confirm for SR26
SR25_CONTACT = {"ABP": "Gerard", "ARTS": "Rose", "FBE": "Katie", "SCI": "Katie",
                "CI&S – Estate Planning & Development": "Gerard and Davina", "CI&S – Sustainability Strategy": "Director",
                "CFOG – Procurement": "Chris", "CFOG – Treasury & Investments": "Chris", "ESG": "All", "CIOG": "TBC",
                "MRE": "Director", "L&R": "Gerard", "SST": "All", "FAC*": "Director", "COO*": "Director"}
CODES = [("A", "Accountable"), ("A and R", "Accountable and Responsible"), ("R", "Responsible – narrative request"),
         ("R (quant)", "Responsible – figures for the Databook (Tab 4)"), ("R (compile)", "SST compiles from other responses / no request"),
         ("S", "Support"), ("C", "Consulted"), ("I", "Informed")]
FILL = {"A and R": "00B050", "A": "92D050", "R": "C6EFCE", "R (quant)": "BDD7EE", "R (compile)": "D9D9D9",
        "S": "FFEB9C", "C": "F8CBAD", "I": "EDEDED"}


def rasci_tab(wb, qual, coverage, faculties, central):
    import build as B
    import quant
    ws = wb.create_sheet("3b. RASCI matrix")
    ws.sheet_view.showGridLines = False
    ws["B1"] = "Sustainability 2030 – SR26 reporting roles (RASCI)"
    ws["B1"].font = Font(name="Aptos", size=20, bold=True, color=B.NAVY)
    ws["B2"] = ("Who is Accountable / Responsible / Support / Consulted / Informed for reporting each Sustainability 2030 indicator in SR26. "
                "Green columns are consolidated portfolios (click [+]/[–] above to show or hide their areas). Codes are pre-filled from the "
                "requests issued; edit with the dropdown. Accountable owners (A) are TBC – to confirm with the Director, Sustainability.")
    ws["B2"].alignment = B.WRAP; ws.merge_cells("B2:G2"); ws.row_dimensions[2].height = 48

    # ---- columns: (portfolio, area label, key, is_consolidated)
    cols = [("Faculty", "Faculties – pilot (all)", "FAC*", True)]
    cols += [("Faculty", f"{code} – {name.replace('Faculty of ', '')}", code, False) for code, name in faculties.items()]
    coo = [(c, n, secs) for c, n, _, secs in central if c != "MRE"]
    cols.append(("COO portfolio", "COO portfolio (all)", "COO*", True))
    for code, name, secs in coo:
        for label, _ in secs:
            key = label.split(" (")[0].split(" – ")[0] if len(secs) == 1 else label.split(" (")[0]
            cols.append(("COO portfolio", label.split(" (")[0], key, False))
    cols.append(("Chancellery", "MRE – Melbourne Research and Enterprise", "MRE", False))
    cols.append(("Sustainability Strategy", "SST (compiles / Databook)", "SST", False))
    first = 9  # first area column (I), as SR25
    col_of = {c[2]: first + i for i, c in enumerate(cols)}

    hdr = ["Number", "Category", "Domain", "Priority area", "Target / indicator", "Type", "Related SR25 ref", "Activity area >>"]
    for i, h in enumerate(hdr):
        ws.cell(6, 1 + i, h)
    ws.cell(3, 8, "SST contact (SR25 – confirm) >>"); ws.cell(4, 8, "Portfolio / Faculty >>")
    ws.merge_cells(start_row=5, start_column=first, end_row=5, end_column=first + len(cols) - 1)
    ws.cell(5, first, "Areas Accountable and/or Responsible for SR26 reporting (green = consolidated portfolio, white = individual area)")
    for i, (port, label, key, cons) in enumerate(cols):
        c = first + i
        ws.cell(3, c, SR25_CONTACT.get(key, SR25_CONTACT.get(key.split(" –")[0], "TBC")))
        ws.cell(4, c, port + (" (consolidated)" if cons else ""))
        ws.cell(6, c, label)
        for r in (3, 4, 6):
            x = ws.cell(r, c); x.alignment = B.CWRAP; x.border = B.BOX
            x.font = Font(name="Aptos", bold=r == 6, size=10 if r != 6 else 11)
        ws.cell(6, c).fill = PatternFill("solid", fgColor="00B050" if cons else "FFFFFF")
        ws.column_dimensions[get_column_letter(c)].width = 13
        if not cons and port in ("Faculty", "COO portfolio"):
            ws.column_dimensions[get_column_letter(c)].outlineLevel = 1
    for i in range(1, 9):
        x = ws.cell(6, i); x.fill = B.HDR; x.font = B.WHITE_B; x.alignment = B.CWRAP; x.border = B.BOX
    ws.cell(5, first).font = Font(name="Aptos", bold=True); ws.cell(5, first).alignment = B.CWRAP
    ws.row_dimensions[6].height = 60

    # ---- who does what per indicator
    cov = {c["ref"]: c for c in coverage}
    purged = quant.quant_only_rows(qual)

    def codes_for(q):
        d = {}
        if "Faculties" in q["stake"]:
            for code in faculties:
                d[code] = "R"
        for code, name, _, secs in central:
            for label, m in secs:
                if m(q):
                    key = "MRE" if code == "MRE" else next(c[2] for c in cols if c[1] == label.split(" (")[0])
                    d[key] = "R"
        c = cov.get(q["ref"])
        if c:
            own = c["owner"]
            if own == "SST":
                d.setdefault("SST", "R (compile)")
            else:
                key = own if own in col_of else next((k for k in col_of if k.startswith(own)), None)
                if key and d.get(key) != "R" and c["npts"]:
                    d[key] = "R (quant)"
                elif key and c["npts"]:
                    d[key] = "R"  # narrative + figures from the same area
        if not d:
            d["SST"] = "R (compile)"
        d["SST"] = d.get("SST", "S")  # SST supports every indicator
        if any(k in faculties for k in d):
            d["FAC*"] = "A and R"
        if any(cols[col_of[k] - first][0] == "COO portfolio" for k in d if k in col_of):
            d["COO*"] = "A and R"
        return d

    r, num = 7, {"dom": 0}
    dom = pa = tgt = None
    data_rows = []
    dfill = PatternFill("solid", fgColor="D9D9D9"); pfill = PatternFill("solid", fgColor="F2F2F2")
    for q in qual:
        if q["domain"] != dom:
            dom = q["domain"]; num["dom"] += 1
            for i, v in enumerate([f"D{num['dom']}", "Domain", "", "", dom]):
                ws.cell(r, 1 + i, v)
            for c in range(1, first + len(cols)):
                ws.cell(r, c).fill = dfill; ws.cell(r, c).font = Font(name="Aptos", bold=True)
            r += 1
        if q["pa"] != pa:
            pa = q["pa"]
            for i, v in enumerate([q["ref"][:2], "Priority", "", "", pa]):
                ws.cell(r, 1 + i, v)
            for c in range(1, first + len(cols)):
                ws.cell(r, c).fill = pfill; ws.cell(r, c).font = Font(name="Aptos", bold=True)
            r += 1
        if q["target"] != tgt:
            tgt = q["target"]
            for i, v in enumerate([q["ref"].split("(")[0], "Target", dom, pa, tgt]):
                c = ws.cell(r, 1 + i, v); c.alignment = B.WRAP; c.font = Font(name="Aptos", bold=True, size=10)
            ws.row_dimensions[r].height = 45
            r += 1
        kind = "Quantitative" if q["row"] in purged else ("Mixed" if q["ref"] in cov else "Qualitative")
        for i, v in enumerate([q["ref"], "Indicator", dom, pa, q["indicator"], kind, B.SR25_MAP.get(q["row"], "")]):
            c = ws.cell(r, 1 + i, v); c.alignment = B.WRAP; c.border = B.BOX; c.font = Font(name="Aptos Narrow", size=10)
        for k, v in codes_for(q).items():
            if k in col_of:
                ws.cell(r, col_of[k], v)
        for c in range(first, first + len(cols)):
            x = ws.cell(r, c); x.alignment = B.CWRAP; x.border = B.BOX; x.font = Font(name="Aptos Narrow", size=10)
        ws.row_dimensions[r].height = 42
        data_rows.append(r)
        r += 1
    top, bot = data_rows[0], data_rows[-1]
    last_col = get_column_letter(first + len(cols) - 1)
    area = f"{get_column_letter(first)}{top}:{last_col}{bot}"
    B.dv_list(ws, '"' + ",".join(c for c, _ in CODES) + '"', area)
    for code, colour in FILL.items():
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'EXACT({get_column_letter(first)}{top},"{code}")'],
                                                        fill=PatternFill("solid", fgColor=colour)))

    # ---- legend + summary (as SR25 rows 50–61)
    r = bot + 2
    ws.cell(r, 5, "Summary").font = Font(name="Aptos", bold=True, size=14, color=B.NAVY)
    r += 1
    ws.cell(r, 5, "Area >>").font = Font(name="Aptos", bold=True)
    for i, c in enumerate(cols):
        ws.cell(r, first + i, f"={get_column_letter(first + i)}6").font = Font(name="Aptos", bold=True, size=9)
        ws.cell(r, first + i).alignment = B.CWRAP
    ws.row_dimensions[r].height = 45
    rows_at = {}
    for code, desc in CODES:
        r += 1
        ws.cell(r, 5, desc); ws.cell(r, 6, code).fill = PatternFill("solid", fgColor=FILL[code])
        rows_at[code] = r
        for i in range(len(cols)):
            L = get_column_letter(first + i)
            ws.cell(r, first + i, f'=COUNTIF({L}{top}:{L}{bot},$F{r})').alignment = B.CWRAP
    for label, f in [("Total A", lambda L: f"={L}{rows_at['A']}+{L}{rows_at['A and R']}"),
                     ("Total R (all types)", lambda L: f"={L}{rows_at['A and R']}+{L}{rows_at['R']}+{L}{rows_at['R (quant)']}+{L}{rows_at['R (compile)']}"),
                     ("Total indicators involved", lambda L: f"=SUM({L}{rows_at['A']}:{L}{rows_at['I']})")]:
        r += 1
        ws.cell(r, 5, label).font = Font(name="Aptos", bold=True)
        for i in range(len(cols)):
            x = ws.cell(r, first + i, f(get_column_letter(first + i))); x.font = Font(name="Aptos", bold=True); x.alignment = B.CWRAP
    B.widths(ws, {"A": 8, "B": 9, "C": 14, "D": 16, "E": 50, "F": 12, "G": 10, "H": 4})
    ws.freeze_panes = f"{get_column_letter(first)}7"
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.auto_filter.ref = f"A6:{last_col}{bot}"
